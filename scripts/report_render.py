"""Render a list of content blocks to .docx (python-docx) and .tex (LaTeX/Overleaf).

Block types
  ("h1", text) ("h2", text) ("p", text) ("bullets", [..]) ("numbered", [..])
  ("table", caption, header, rows) ("figure", path, caption, width_in) ("code", text) ("pagebreak",)
Inline **bold**, [label](url) and bare https:// URLs are supported in paragraphs, bullets and table cells.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

NAVY = RGBColor(0x0B, 0x2A, 0x4A)
BLUE = RGBColor(0x2F, 0x6F, 0xAE)
BOLD_RE = re.compile(r"(\*\*[^*]+\*\*)")


# ------------------------------------------------------------------ DOCX helpers
def _field(paragraph, instr: str):
    r = paragraph.add_run()
    for tag, attr in [("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", None),
                      ("w:fldChar", {"w:fldCharType": "separate"}), ("w:t", None), ("w:fldChar", {"w:fldCharType": "end"})]:
        el = OxmlElement(tag)
        if attr:
            for k, v in attr.items():
                el.set(qn(k), v)
        if tag == "w:instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        if tag == "w:t":
            el.text = "Right-click and choose Update Field to refresh." if "TOC" in instr else "1"
        r._r.append(el)


# [label](https://...) or a bare https://... URL (trailing punctuation is not part of the link)
LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)|(https?://[^\s)]*[^\s).,;:])")
HYPERLINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"


def _hyperlink(paragraph, url, text, size=None, bold=False):
    """Append a clickable external hyperlink (blue, underlined) to a paragraph."""
    rid = paragraph.part.relate_to(url, HYPERLINK_REL, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if bold:
        rPr.append(OxmlElement("w:b"))
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "1F5FA8")
    rPr.append(c)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(round(size * 2))))
        rPr.append(sz)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def _plain(paragraph, text, bold, size, color):
    run = paragraph.add_run(text)
    run.bold = bold or None
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def _runs(paragraph, text, size=None, color=None):
    """Inline **bold**, [label](url) and bare URLs (rendered as clickable links)."""
    for part in BOLD_RE.split(text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        body = part[2:-2] if bold else part
        pos = 0
        for m in LINK_RE.finditer(body):
            if m.start() > pos:
                _plain(paragraph, body[pos:m.start()], bold, size, color)
            label, url = (m.group(1), m.group(2)) if m.group(2) else (m.group(3), m.group(3))
            _hyperlink(paragraph, url, label, size, bold)
            pos = m.end()
        if pos < len(body):
            _plain(paragraph, body[pos:], bold, size, color)


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _keep_table_together(t, small: bool):
    """Repeat the header row on every page, never split a row, and keep short tables on one page."""
    for i, r in enumerate(t.rows):
        trPr = r._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
        if i == 0:
            trPr.append(OxmlElement("w:tblHeader"))
        if small and i < len(t.rows) - 1:
            for c in r.cells:
                for para in c.paragraphs:
                    para.paragraph_format.keep_with_next = True


def _setup(doc):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ["left_margin", "right_margin", "top_margin", "bottom_margin"]:
        setattr(sec, side, Cm(2.4))
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for name, size in [("Heading 1", 16), ("Heading 2", 13), ("Title", 24)]:
        s = doc.styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.color.rgb = NAVY
        s.font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(18)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    cap = doc.styles["Caption"]
    cap.font.size = Pt(9.5)
    cap.font.italic = False
    cap.font.bold = False
    cap.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    # footer page number
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ")
    _field(fp, "PAGE")
    sec.different_first_page_header_footer = True


def title_page(doc, meta):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(p, meta["kicker"], size=12, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta["title"])
    r.bold, r.font.size, r.font.color.rgb = True, Pt(22), NAVY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta["subtitle"])
    r.font.size, r.font.color.rgb = Pt(14), BLUE
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta["doc_type"])
    r.bold, r.font.size = True, Pt(13)
    for _ in range(3):
        doc.add_paragraph()
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in meta["fields"]:
        c = t.add_row().cells
        c[0].text = k
        c[0].paragraphs[0].runs[0].bold = True
        _runs(c[1].paragraphs[0], v)
        c[0].width, c[1].width = Cm(5), Cm(9)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(p, meta["note"], size=9.5, color=RGBColor(0x5F, 0x6B, 0x76))
    p.add_run().add_break(WD_BREAK.PAGE)


def to_docx(blocks, meta, out: Path):
    doc = Document()
    _setup(doc)
    title_page(doc, meta)
    doc.add_paragraph("Table of Contents", style="Heading 1")
    _field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    for b in blocks:
        kind = b[0]
        if kind in ("h1", "h2"):
            doc.add_paragraph(b[1], style="Heading 1" if kind == "h1" else "Heading 2")
        elif kind == "p":
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _runs(p, b[1])
        elif kind in ("bullets", "numbered"):
            for it in b[1]:
                p = doc.add_paragraph(style="List Bullet" if kind == "bullets" else "List Number")
                _runs(p, it)
        elif kind == "table":
            _, caption, header, rows = b[:4]
            cp = doc.add_paragraph(style="Caption")
            _runs(cp, caption)
            cp.paragraph_format.keep_with_next = True
            t = doc.add_table(rows=1 + len(rows), cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(header):
                c = t.rows[0].cells[j]
                c.text = ""
                _runs(c.paragraphs[0], f"**{h}**", size=9.5, color=RGBColor(0xFF, 0xFF, 0xFF))
                _shade(c, "2F6FAE")
            for i, r in enumerate(rows, 1):
                for j, v in enumerate(r):
                    c = t.rows[i].cells[j]
                    c.text = ""
                    _runs(c.paragraphs[0], str(v), size=9.5)
                    if i % 2 == 0:
                        _shade(c, "EEF4FA")
            _keep_table_together(t, small=len(rows) <= 14)
            doc.add_paragraph()
        elif kind == "figure":
            _, path, caption, width = b
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Inches(width))
            cp = doc.add_paragraph(style="Caption")
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(cp, caption)
        elif kind == "code":
            p = doc.add_paragraph()
            r = p.add_run(b[1])
            r.font.name, r.font.size = "Consolas", Pt(8.5)
            p.paragraph_format.left_indent = Cm(0.5)
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.core_properties.title = meta["title"]
    doc.core_properties.author = meta.get("author", "")
    doc.save(out)


# ------------------------------------------------------------------ LaTeX
TEX_ESC = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_", "{": r"\{",
           "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
UNI = {"≤": r"$\leq$", "≥": r"$\geq$", "±": r"$\pm$", "×": r"$\times$", "²": r"$^2$", "°": r"$^\circ$", "→": r"$\rightarrow$",
       "–": "--", "—": "---", "’": "'", "“": "``", "”": "''", "≈": r"$\approx$", "ε": r"$\varepsilon$", "λ": r"$\lambda$",
       "−": "-", "·": r"$\cdot$", "≠": r"$\neq$", "↔": r"$\leftrightarrow$", "…": r"\ldots{}", "κ": r"$\kappa$", "∈": r"$\in$"}


def _esc_plain(s: str) -> str:
    out = "".join(TEX_ESC.get(ch, ch) for ch in s)
    for k, v in UNI.items():
        out = out.replace(k, v)
    return out


def esc(s: str) -> str:
    """Escape text for LaTeX: [label](url) -> \\href, bare URLs -> \\url, **bold** -> \\textbf."""
    out, pos = [], 0
    for m in LINK_RE.finditer(s):
        out.append(_esc_plain(s[pos:m.start()]))
        if m.group(2):
            out.append(r"\href{" + m.group(2).replace("%", r"\%").replace("#", r"\#") + "}{" + _esc_plain(m.group(1)) + "}")
        else:
            out.append(r"\url{" + m.group(3) + "}")
        pos = m.end()
    out.append(_esc_plain(s[pos:]))
    return re.sub(r"\*\*([^*]+)\*\*", r"\\textbf{\1}", "".join(out))


def to_tex(blocks, meta, out: Path, fig_prefix="../results/figures/"):
    L = [r"\documentclass[11pt,a4paper]{article}",
         r"\usepackage[margin=2.4cm]{geometry}", r"\usepackage[T1]{fontenc}", r"\usepackage{lmodern}",
         r"\usepackage{graphicx}", r"\usepackage{booktabs}", r"\usepackage{tabularx}", r"\usepackage{caption}",
         r"\usepackage{xcolor}", r"\usepackage[hidelinks]{hyperref}", r"\usepackage{enumitem}", r"\usepackage{float}",
         r"\usepackage{fancyhdr}", r"\pagestyle{fancy}\fancyhf{}\cfoot{Page \thepage}\renewcommand{\headrulewidth}{0pt}",
         r"\definecolor{navy}{HTML}{0B2A4A}", r"\setlength{\parskip}{4pt}",
         r"\begin{document}", r"\begin{titlepage}\centering\vspace*{2cm}",
         r"{\color{navy}\large " + esc(meta["kicker"]) + r"\par}\vspace{0.8cm}",
         r"{\color{navy}\LARGE\bfseries " + esc(meta["title"]) + r"\par}\vspace{0.5cm}",
         r"{\large " + esc(meta["subtitle"]) + r"\par}\vspace{1cm}",
         r"{\bfseries " + esc(meta["doc_type"]) + r"\par}\vspace{2cm}",
         r"\begin{tabular}{ll}"]
    L += [r"\textbf{" + esc(k) + "} & " + esc(v) + r" \\" for k, v in meta["fields"]]
    L += [r"\end{tabular}\vfill", r"{\small " + esc(meta["note"]) + r"\par}", r"\end{titlepage}",
          r"\tableofcontents\newpage"]
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            L.append(r"\section*{" + esc(b[1]) + r"}\addcontentsline{toc}{section}{" + esc(b[1]) + "}")
        elif kind == "h2":
            L.append(r"\subsection*{" + esc(b[1]) + r"}\addcontentsline{toc}{subsection}{" + esc(b[1]) + "}")
        elif kind == "p":
            L.append(esc(b[1]) + "\n")
        elif kind in ("bullets", "numbered"):
            env = "itemize" if kind == "bullets" else "enumerate"
            L.append(r"\begin{" + env + r"}[leftmargin=*,itemsep=1pt]")
            L += [r"\item " + esc(it) for it in b[1]]
            L.append(r"\end{" + env + "}")
        elif kind == "table":
            _, caption, header, rows = b[:4]
            n = len(header)
            spec = "X" * n
            L += [r"\begin{table}[H]\centering\small", r"\caption*{" + esc(caption) + "}",
                  r"\begin{tabularx}{\linewidth}{" + spec + "}", r"\toprule",
                  " & ".join(r"\textbf{" + esc(h) + "}" for h in header) + r" \\", r"\midrule"]
            L += [" & ".join(esc(str(v)) for v in r) + r" \\" for r in rows]
            L += [r"\bottomrule", r"\end{tabularx}", r"\end{table}"]
        elif kind == "figure":
            _, path, caption, width = b
            L += [r"\begin{figure}[H]\centering",
                  r"\includegraphics[width=" + f"{min(width / 6.3, 1.0):.2f}" + r"\linewidth]{" + fig_prefix + Path(path).name + "}",
                  r"\caption*{" + esc(caption) + "}", r"\end{figure}"]
        elif kind == "code":
            L += [r"\begin{verbatim}", b[1], r"\end{verbatim}"]
        elif kind == "pagebreak":
            L.append(r"\newpage")
    L.append(r"\end{document}")
    out.write_text("\n".join(L), encoding="utf-8")
